from pptx import Presentation

prs = Presentation("file_cua_ban.pptx")
md_content = []

for i, slide in enumerate(prs.slides, start=1):
    md_content.append(f"\n# Slide {i}\n")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text:
                    md_content.append(f"- {text}")

with open("output.md", "w", encoding="utf-8") as f:
    f.write("\n".join(md_content))

print("Đã chuyển đổi thành công sang output.md!")